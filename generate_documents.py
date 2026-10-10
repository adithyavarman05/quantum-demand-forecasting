"""
Generate Capstone_PPT.pptx and Capstone_Report.docx
Uses ONLY real metrics from artifacts.joblib — no invented numbers.
"""
import os
import joblib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
art = joblib.load(os.path.join(BASE_DIR, "artifacts.joblib"))

# ---- Extract real metrics ----
mae_vqc, rmse_vqc, r2_vqc = art["metrics_vqc"]
mae_lr, rmse_lr, r2_lr = art["metrics_lr"]
vqc_time = art["vqc_time"]
lr_time = art["lr_time"]
n_params_vqc = art["n_params_vqc"]
n_params_lr = art["n_params_lr"]
final_loss = art["losses"][-1]

# Colors for dark theme
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)
BG_SLIDE = RGBColor(0x16, 0x21, 0x3E)
TEXT_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_LIGHT = RGBColor(0xD0, 0xD0, 0xE0)
ACCENT_PURPLE = RGBColor(0x7C, 0x4D, 0xFF)
ACCENT_BLUE = RGBColor(0x66, 0x7E, 0xEA)
ACCENT_CYAN = RGBColor(0x00, 0xD2, 0xFF)


# ============================================================
# PPTX GENERATION
# ============================================================
def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_text(slide, text, left, top, width, height, font_size=28, color=TEXT_WHITE, bold=True, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return tf


def add_bullets(slide, items, left, top, width, height, font_size=16, color=TEXT_LIGHT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.level = 0
    return tf


def generate_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # blank

    L = Inches(0.8)
    W = Inches(11.5)

    # ===== SLIDE 1: Title =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_DARK)
    add_title_text(slide, "⚡ Quantum AI-based Electricity\nDemand Forecasting System",
                   L, Inches(1.5), W, Inches(2.5), font_size=40, color=ACCENT_CYAN, alignment=PP_ALIGN.CENTER)
    add_title_text(slide, "Hybrid Variational Quantum Circuit (VQC) vs Linear Regression",
                   L, Inches(3.8), W, Inches(0.8), font_size=22, color=TEXT_LIGHT, bold=False, alignment=PP_ALIGN.CENTER)
    add_title_text(slide, "Adithya Varman Nandhivarman",
                   L, Inches(4.8), W, Inches(0.6), font_size=20, color=TEXT_WHITE, alignment=PP_ALIGN.CENTER)
    add_title_text(slide, "Course: 23CSE304  •  Quantum Computing Capstone Project",
                   L, Inches(5.4), W, Inches(0.6), font_size=16, color=TEXT_LIGHT, bold=False, alignment=PP_ALIGN.CENTER)

    # ===== SLIDE 2: Problem Statement =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Problem Statement", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    add_bullets(slide, [
        "Accurate short-term electricity demand forecasting is critical for grid stability and cost optimization.",
        "Classical models (e.g., Linear Regression) are fast but may miss complex nonlinear patterns.",
        "Can hybrid quantum-classical machine learning models improve forecasting accuracy?",
        "This project investigates a Variational Quantum Circuit (VQC) for hourly electricity demand prediction.",
        "Dataset: AEP hourly electricity consumption from PJM Interconnection (Kaggle).",
    ], L, Inches(1.5), W, Inches(5.5))

    # ===== SLIDE 3: Objectives =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Objectives", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    add_bullets(slide, [
        "Build a hybrid quantum-classical regression model using PennyLane's Variational Quantum Circuit.",
        "Compare VQC performance against a classical Linear Regression baseline.",
        "Evaluate using MAE, RMSE, and R² on real-world hourly electricity demand data.",
        "Deploy an interactive Streamlit web app for live predictions and result visualization.",
        "Provide honest analysis: report whether quantum advantage is observed at this scale.",
    ], L, Inches(1.5), W, Inches(5.5))

    # ===== SLIDE 4: Proposed Algorithm (Explanation) =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Proposed Algorithm: Hybrid VQC Regressor", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    add_bullets(slide, [
        "Step 1: Create 4 lag features from hourly demand (t-4, t-3, t-2, t-1) → predict t.",
        "Step 2: Scale features to [0,1] and target to [-1,1] (MinMaxScaler, fit on train only).",
        "Step 3: Encode inputs via AngleEmbedding (RX gates, angle = π × x) on 4 qubits.",
        "Step 4: Apply StronglyEntanglingLayers (2 layers, 24 params: Rot + CNOT entanglement).",
        "Step 5: Measure ⟨Z₀⟩ as output, optimize MSE loss with Adam (lr=0.05, 80 epochs).",
    ], L, Inches(1.5), W, Inches(5.5))

    # ===== SLIDE 5: Pseudocode =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Proposed Algorithm: Pseudocode", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    pseudocode = [
        "1. Load AEP_hourly.csv, sort by Datetime, take first 1500 rows",
        "2. Create lag matrix X[i] = [demand[i-4], ..., demand[i-1]], y[i] = demand[i]",
        "3. Split 80/20 chronologically → X_train, X_test, y_train, y_test",
        "4. Fit MinMaxScaler on train; transform both train and test",
        "5. Initialize VQC weights θ ~ Uniform(0, 2π), shape = (2, 4, 3)",
        "6. FOR epoch = 1 to 80:",
        "      loss = MSE( VQC(X_train, θ), y_train )",
        "      θ ← Adam_update(θ, ∇loss)     # backprop on simulator",
        "7. Evaluate VQC and Linear Regression on test set in original MW units",
    ]
    add_bullets(slide, pseudocode, L, Inches(1.5), W, Inches(5.5), font_size=15)

    # ===== SLIDE 6: Algorithm Complexity =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Algorithm Complexity", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    add_bullets(slide, [
        f"Time Complexity (Simulator): O(E × n × L × N × 2ⁿ)",
        f"  E={80} epochs, n={4} qubits, L={2} layers, N={1196} samples, 2ⁿ={16} statevector",
        f"Space Complexity: O(2ⁿ + N×n + n×L) = O(16 + {1196}×4 + 4×2)",
        "On real quantum hardware, gradients use the parameter-shift rule (2 circuit evaluations per parameter per gradient).",
        "This project uses backpropagation on the PennyLane default.qubit simulator for speed.",
    ], L, Inches(1.5), W, Inches(5.5))

    # ===== SLIDE 7: Implementation =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Implementation", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    add_bullets(slide, [
        "Language: Python 3.14  |  Quantum Framework: PennyLane 0.45",
        "Libraries: NumPy, Pandas, scikit-learn, Matplotlib, Streamlit, Joblib",
        "Simulator: default.qubit (statevector, diff_method='backprop')",
        "Web App: Streamlit with 3 tabs (Predict, Results, Circuit)",
    ], L, Inches(1.5), W, Inches(2.5))
    # Code snippet
    code = (
        "@qml.qnode(dev, interface='autograd', diff_method='backprop')\n"
        "def vqc(x, weights):\n"
        "    qml.AngleEmbedding(π*x, wires=range(4), rotation='X')\n"
        "    qml.StronglyEntanglingLayers(weights, wires=range(4))\n"
        "    return qml.expval(qml.PauliZ(0))"
    )
    add_title_text(slide, code, Inches(1.5), Inches(4.0), Inches(10), Inches(2.5),
                   font_size=14, color=ACCENT_CYAN, bold=False)

    # ===== SLIDE 8: Results and Analysis =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Results and Analysis", L, Inches(0.3), W, Inches(0.6), font_size=32, color=ACCENT_CYAN)

    # Comparison table as text
    table_text = (
        f"{'Metric':<18}{'VQC (Quantum)':>16}{'Linear Regression':>20}\n"
        f"{'─' * 54}\n"
        f"{'MAE (MW)':<18}{mae_vqc:>16.4f}{mae_lr:>20.4f}\n"
        f"{'RMSE (MW)':<18}{rmse_vqc:>16.4f}{rmse_lr:>20.4f}\n"
        f"{'R²':<18}{r2_vqc:>16.4f}{r2_lr:>20.4f}\n"
        f"{'Train time (s)':<18}{vqc_time:>16.2f}{lr_time:>20.4f}\n"
        f"{'Parameters':<18}{n_params_vqc:>16d}{n_params_lr:>20d}"
    )
    add_title_text(slide, table_text, Inches(0.5), Inches(1.1), Inches(6.5), Inches(3.0),
                   font_size=13, color=TEXT_LIGHT, bold=False)

    # Add results plot image
    results_img = os.path.join(BASE_DIR, "results_plots.png")
    if os.path.exists(results_img):
        slide.shapes.add_picture(results_img, Inches(0.5), Inches(4.2), Inches(8.5), Inches(3.0))

    # Add screenshot
    ss_predict = os.path.join(BASE_DIR, "screenshot_predict.png")
    if os.path.exists(ss_predict):
        slide.shapes.add_picture(ss_predict, Inches(9.2), Inches(1.1), Inches(3.8), Inches(2.8))

    ss_results = os.path.join(BASE_DIR, "screenshot_results.png")
    if os.path.exists(ss_results):
        slide.shapes.add_picture(ss_results, Inches(9.2), Inches(4.2), Inches(3.8), Inches(2.8))

    # ===== SLIDE 9: Conclusion =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_SLIDE)
    add_title_text(slide, "Conclusion", L, Inches(0.4), W, Inches(0.8), font_size=32, color=ACCENT_CYAN)
    add_bullets(slide, [
        f"Linear Regression achieves R² = {r2_lr:.4f}, significantly outperforming VQC (R² = {r2_vqc:.4f}).",
        "This is expected: with only 4 qubits, 24 parameters, and mostly linear data, the VQC cannot capture the full pattern.",
        "The VQC's expressivity is limited by the small Hilbert space (2⁴ = 16 dimensions) on a classical simulator.",
        "No quantum advantage is claimed. On real hardware, noise and decoherence would further degrade performance.",
        "Value: This project demonstrates a working hybrid quantum-classical ML pipeline, ready for scaling when hardware matures.",
    ], L, Inches(1.5), W, Inches(5.5))

    # ===== SLIDE 10: Thank You =====
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_DARK)
    add_title_text(slide, "Thank You", L, Inches(2.0), W, Inches(1.5), font_size=48, color=ACCENT_CYAN, alignment=PP_ALIGN.CENTER)
    add_title_text(slide, "Adithya Varman Nandhivarman  •  23CSE304\nQuantum AI-based Electricity Demand Forecasting",
                   L, Inches(3.8), W, Inches(1.5), font_size=20, color=TEXT_LIGHT, bold=False, alignment=PP_ALIGN.CENTER)
    add_title_text(slide, "References:\n• PennyLane Documentation — https://docs.pennylane.ai\n• AEP Hourly Dataset — Kaggle (PJM Interconnection)",
                   L, Inches(5.5), W, Inches(1.5), font_size=14, color=TEXT_LIGHT, bold=False, alignment=PP_ALIGN.CENTER)

    pptx_path = os.path.join(BASE_DIR, "Capstone_PPT.pptx")
    prs.save(pptx_path)
    print(f"Saved: {pptx_path}")


# ============================================================
# DOCX GENERATION
# ============================================================
def generate_docx():
    doc = Document()

    # ---- Styles ----
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = DocPt(11)
    style.paragraph_format.space_after = DocPt(6)

    def add_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = DocRGBColor(0x1A, 0x23, 0x7E)
        return h

    def add_para(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.alignment = align
        run = p.add_run(text)
        run.bold = bold
        return p

    def add_image(filename, width=DocInches(5.5)):
        path = os.path.join(BASE_DIR, filename)
        if os.path.exists(path):
            doc.add_picture(path, width=width)
            last_paragraph = doc.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ===== TITLE PAGE =====
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Quantum AI-based Electricity\nDemand Forecasting System")
    run.font.size = DocPt(26)
    run.bold = True
    run.font.color.rgb = DocRGBColor(0x1A, 0x23, 0x7E)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Hybrid Variational Quantum Circuit (VQC) vs Linear Regression")
    run.font.size = DocPt(14)
    run.font.color.rgb = DocRGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()
    for line in [
        "Adithya Varman Nandhivarman",
        "Course: 23CSE304 — Quantum Computing",
        "Capstone Project Report",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.size = DocPt(13)

    doc.add_page_break()

    # ===== ABSTRACT =====
    add_heading("Abstract")
    add_para(
        "This report presents a hybrid quantum-classical machine learning approach to short-term electricity demand forecasting. "
        "A Variational Quantum Circuit (VQC) with 4 qubits and 2 StronglyEntanglingLayers (24 trainable parameters) is implemented "
        "using PennyLane's default.qubit simulator with backpropagation-based gradient computation. The model is trained on 1,500 hourly "
        "demand readings from the AEP (American Electric Power) dataset obtained from Kaggle. Performance is compared against a classical "
        f"Linear Regression baseline. The Linear Regression model achieves an R² of {r2_lr:.4f} with MAE of {mae_lr:.4f} MW, significantly "
        f"outperforming the VQC (R² = {r2_vqc:.4f}, MAE = {mae_vqc:.4f} MW). This result is expected given the small circuit size, limited "
        "expressivity of 4 qubits, and the predominantly linear nature of the lag-based features. No quantum advantage is claimed. "
        "The project demonstrates a complete, working hybrid quantum-classical pipeline including model training, evaluation, and an interactive "
        "Streamlit web application for live predictions."
    )

    # ===== PROBLEM STATEMENT =====
    add_heading("Problem Statement")

    add_heading("Problem", level=2)
    add_para(
        "Electricity demand forecasting is essential for power grid operators to balance supply and demand, minimize costs, "
        "and prevent blackouts. Short-term (hourly) forecasting is particularly challenging due to rapid fluctuations driven "
        "by weather, time of day, and human behaviour. While classical statistical and machine learning methods perform well, "
        "quantum computing promises potential advantages for certain optimization and learning tasks. This project investigates "
        "whether a hybrid Variational Quantum Circuit can provide competitive forecasting performance."
    )

    add_heading("Input", level=2)
    add_para(
        "The input is the AEP_hourly.csv dataset from the PJM Interconnection energy market on Kaggle. It contains hourly "
        "electricity consumption in megawatts (MW) with columns: Datetime and AEP_MW. The first 1,500 rows (sorted chronologically) "
        "are used. Four lag features are constructed: the demand values at hours t-4, t-3, t-2, and t-1, used to predict the demand at hour t."
    )

    add_heading("Expected Output", level=2)
    add_para(
        "The system outputs a predicted electricity demand value (in MW) for the next hour, given the previous four hours' demand. "
        "Both the VQC and Linear Regression models produce these predictions, which are compared against actual values using MAE, RMSE, and R² metrics."
    )

    add_heading("Limitations", level=2)
    add_para(
        "1. The VQC uses only 4 qubits and 24 parameters, severely limiting its expressivity compared to classical models.\n"
        "2. The default.qubit simulator runs on a classical computer, offering no actual quantum speedup.\n"
        "3. Only 4 lag features are used; richer feature engineering (weather, calendar features) would improve both models.\n"
        "4. On real quantum hardware, gate noise and decoherence would further degrade VQC performance.\n"
        "5. The dataset is limited to 1,500 rows for computational feasibility on the simulator."
    )

    # ===== OBJECTIVES =====
    add_heading("Objectives")
    objectives = [
        "Implement a hybrid Variational Quantum Circuit (VQC) regressor using PennyLane.",
        "Train the VQC on real-world hourly electricity demand data from the AEP dataset.",
        "Compare VQC performance against a classical Linear Regression baseline using MAE, RMSE, and R².",
        "Deploy an interactive Streamlit web application for live demand predictions.",
        "Provide an honest, evidence-based analysis of quantum vs. classical performance at this scale.",
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')

    # ===== PROPOSED ALGORITHM =====
    add_heading("Proposed Algorithm")

    add_heading("Algorithm Explanation", level=2)
    add_para(
        "The proposed algorithm is a Hybrid Variational Quantum Circuit (VQC) Regressor. It combines a parameterized quantum circuit "
        "with classical optimization to perform regression. The quantum circuit acts as a trainable function: input features are encoded "
        "into qubit rotations (AngleEmbedding with RX gates), a parameterized ansatz (StronglyEntanglingLayers) creates entanglement "
        "and applies learnable rotations, and the expectation value of the Pauli-Z operator on qubit 0 serves as the scalar output. "
        "A classical optimizer (Adam) minimizes the mean squared error between the circuit output and the scaled target values."
    )

    add_heading("Flowchart", level=2)
    add_image("quantum_circuit.png", width=DocInches(5.5))
    add_para("Figure: 4-qubit Variational Quantum Circuit with AngleEmbedding and 2 StronglyEntanglingLayers.", bold=True)

    add_heading("Step-by-Step Process", level=2)
    steps = [
        "Load the AEP_hourly.csv dataset and sort by Datetime. Use the first 1,500 rows.",
        "Construct 4 lag features: X[i] = [demand[i-4], demand[i-3], demand[i-2], demand[i-1]], y[i] = demand[i].",
        "Split data 80/20 chronologically (no shuffling): 1,196 training samples, 300 test samples.",
        "Fit MinMaxScaler on training data: features scaled to [0, 1], target scaled to [-1, 1].",
        "Initialize VQC weights θ randomly from Uniform(0, 2π) with shape (2, 4, 3) = 24 parameters.",
        "For each of 80 epochs: compute VQC predictions via parameter broadcasting, calculate MSE loss, update θ with Adam (lr = 0.05).",
        "Train Linear Regression baseline on the same scaled training data.",
        "Evaluate both models on the test set, inverse-transforming predictions to original MW units.",
        "Compute MAE, RMSE, and R² for both models and generate comparison plots.",
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"Step {i}: {step}", style='List Number')

    add_heading("Applications", level=2)
    add_para(
        "Hybrid quantum-classical models have potential applications in: (1) energy grid demand forecasting and load balancing, "
        "(2) financial time-series prediction, (3) supply chain demand planning, and (4) any regression task where quantum feature "
        "spaces might capture complex correlations. However, current NISQ-era devices are limited in qubit count and coherence time, "
        "making practical advantage unlikely for problems where classical models already perform well."
    )

    # ===== PSEUDOCODE =====
    add_heading("Pseudocode")
    pseudocode = """ALGORITHM: Hybrid VQC Demand Forecasting

INPUT: Hourly demand series D[1..1500]
OUTPUT: Predicted demand for next hour

1.  FOR i = 5 TO 1500:
        X[i] ← [D[i-4], D[i-3], D[i-2], D[i-1]]
        y[i] ← D[i]

2.  split ← 0.8 × |X|
    X_train, y_train ← X[1..split], y[1..split]
    X_test, y_test   ← X[split+1..|X|], y[split+1..|y|]

3.  x_scaler ← fit MinMaxScaler(0,1) on X_train
    y_scaler ← fit MinMaxScaler(-1,1) on y_train
    X_train ← x_scaler.transform(X_train)
    X_test  ← clip(x_scaler.transform(X_test), 0, 1)
    y_train ← y_scaler.transform(y_train)

4.  θ ← Random(0, 2π, shape=(2, 4, 3))   // 24 parameters

5.  FOR epoch = 1 TO 80:
        // Quantum circuit with parameter broadcasting
        ŷ ← VQC(X_train, θ)   // AngleEmbedding + StronglyEntanglingLayers + ⟨Z₀⟩
        loss ← MSE(ŷ, y_train)
        θ ← Adam_update(θ, ∇loss, lr=0.05)

6.  pred_vqc  ← y_scaler.inverse(VQC(X_test, θ))
    pred_lr   ← y_scaler.inverse(LinearRegression.predict(X_test))

7.  RETURN MAE, RMSE, R² for both models"""

    p = doc.add_paragraph()
    run = p.add_run(pseudocode)
    run.font.name = 'Consolas'
    run.font.size = DocPt(9)

    # ===== IMPLEMENTATION =====
    add_heading("Implementation")

    add_heading("Language and Tools", level=2)
    add_para(
        "The project is implemented in Python 3.14. The quantum circuit is built using PennyLane 0.45, "
        "an open-source quantum machine learning framework by Xanadu. The classical baseline uses scikit-learn's "
        "LinearRegression. Data processing uses NumPy and Pandas. Visualization uses Matplotlib. The web application "
        "is built with Streamlit. Model artifacts are serialized with Joblib."
    )

    add_heading("Code Explanation", level=2)
    add_para(
        "The core quantum circuit is defined as a PennyLane QNode with the default.qubit device (statevector simulator). "
        "The diff_method='backprop' flag enables efficient gradient computation by backpropagating through the simulated "
        "quantum circuit — this is faster than the parameter-shift rule but only works on simulators. On real quantum hardware, "
        "the parameter-shift rule would be used, requiring 2 circuit evaluations per parameter per gradient step.\n\n"
        "AngleEmbedding encodes each of the 4 input features as an RX rotation (angle = π × x) on one qubit. "
        "StronglyEntanglingLayers applies 2 layers of parameterized Rot gates followed by CNOT entanglement, "
        "creating a 24-parameter trainable ansatz. The circuit output is the expectation value ⟨Z₀⟩ ∈ [-1, 1], "
        "which is mapped back to MW units via inverse scaling.\n\n"
        "Parameter broadcasting allows the entire training batch to be processed in a single circuit call, "
        "significantly improving performance on the simulator."
    )

    code_snippet = '''@qml.qnode(dev, interface="autograd", diff_method="backprop")
def vqc(x, weights):
    qml.AngleEmbedding(np.pi * x, wires=range(4), rotation="X")
    qml.StronglyEntanglingLayers(weights, wires=range(4))
    return qml.expval(qml.PauliZ(0))'''
    p = doc.add_paragraph()
    run = p.add_run(code_snippet)
    run.font.name = 'Consolas'
    run.font.size = DocPt(9)

    # ===== WEBSITE =====
    add_heading("Website (Streamlit Application)")

    add_heading("Technologies and Tools", level=2)
    add_para(
        "The web application is built with Streamlit 1.65, a Python framework for creating interactive data applications. "
        "It requires no frontend JavaScript code — the entire app is written in Python. The app loads pre-trained model "
        "artifacts (weights, scalers, Linear Regression model) from artifacts.joblib and recreates the quantum circuit "
        "at inference time using PennyLane."
    )

    add_heading("Design and Features", level=2)
    add_para(
        "The application features a dark-themed UI with glassmorphism-styled metric cards and three tabs:\n\n"
        "1. Predict Tab: Users enter four hourly demand values (MW) as lag features. Clicking 'Predict Next Hour' runs "
        "both the VQC and Linear Regression models, displaying predictions in styled metric cards.\n\n"
        "2. Results Tab: Displays a comparison table of MAE, RMSE, R², training time, and parameter count for both models, "
        "along with the generated results plots (actual vs. predicted, loss curve, model comparison).\n\n"
        "3. Circuit Tab: Shows the quantum circuit diagram generated by qml.draw_mpl, with an explanation of the architecture."
    )

    add_heading("Screenshots", level=2)
    add_image("screenshot_predict.png", width=DocInches(5.5))
    add_para("Figure: Predict Tab — Enter 4 lag values and get predictions from both models.", bold=True)

    doc.add_paragraph()
    add_image("screenshot_results.png", width=DocInches(5.5))
    add_para("Figure: Results Tab — Performance comparison table and visualization plots.", bold=True)

    doc.add_paragraph()
    add_image("screenshot_circuit.png", width=DocInches(5.5))
    add_para("Figure: Circuit Tab — Quantum circuit architecture diagram.", bold=True)

    # ===== RESULTS =====
    add_heading("Results")

    add_heading("Input", level=2)
    add_para(
        "The input dataset consists of 1,500 hourly electricity demand readings from the AEP region, sorted chronologically. "
        "Four lag features are constructed from each consecutive window of 4 hours, producing 1,496 samples. "
        "These are split 80/20 into 1,196 training and 300 test samples. Features are scaled to [0, 1] and targets to [-1, 1]."
    )

    add_heading("Output and Performance", level=2)
    # Create a proper table
    table = doc.add_table(rows=6, cols=3, style='Light Shading Accent 1')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Metric", "VQC (Quantum)", "Linear Regression"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.bold = True

    data = [
        ("MAE (MW)", f"{mae_vqc:.4f}", f"{mae_lr:.4f}"),
        ("RMSE (MW)", f"{rmse_vqc:.4f}", f"{rmse_lr:.4f}"),
        ("R²", f"{r2_vqc:.4f}", f"{r2_lr:.4f}"),
        ("Train Time (s)", f"{vqc_time:.2f}", f"{lr_time:.4f}"),
        ("Parameters", str(n_params_vqc), str(n_params_lr)),
    ]
    for row_idx, (metric, vqc_val, lr_val) in enumerate(data, 1):
        table.rows[row_idx].cells[0].text = metric
        table.rows[row_idx].cells[1].text = vqc_val
        table.rows[row_idx].cells[2].text = lr_val

    doc.add_paragraph()
    add_image("results_plots.png", width=DocInches(6.0))
    add_para("Figure: (Left) Actual vs. VQC Predicted, (Center) Training Loss Curve, (Right) Model Comparison.", bold=True)

    add_heading("Analysis", level=2)
    add_para(
        f"Linear Regression decisively outperforms the VQC across all metrics. With an R² of {r2_lr:.4f} vs. {r2_vqc:.4f}, "
        f"the classical model explains {r2_lr*100:.1f}% of the variance compared to {r2_vqc*100:.1f}% for the VQC. "
        f"The VQC's MAE ({mae_vqc:.1f} MW) is {mae_vqc/mae_lr:.1f}× worse than Linear Regression ({mae_lr:.1f} MW).\n\n"
        "This outcome is expected for several reasons:\n"
        "1. The 4-qubit circuit operates in a 16-dimensional Hilbert space with only 24 parameters — insufficient to capture "
        "the full complexity of the demand signal.\n"
        "2. Lag-based demand features are predominantly linear (autocorrelated), favouring Linear Regression.\n"
        "3. The VQC output is bounded to [-1, 1] (single Pauli-Z expectation), requiring the entire demand range to be compressed "
        "into this interval, which can lose resolution.\n"
        f"4. Despite 80 epochs of training ({vqc_time:.2f}s), the VQC loss plateaus at {final_loss:.6f}, suggesting the "
        "circuit has reached its representational capacity.\n\n"
        "No quantum advantage or 'exponential speedup' is observed or claimed. The simulator runs on classical hardware, "
        "and the VQC's computational cost exceeds that of Linear Regression by orders of magnitude."
    )

    add_heading("Time and Space Complexity", level=2)
    add_para(
        "Time Complexity on Simulator: O(E × n × L × N × 2ⁿ)\n"
        f"Where E = {80} (epochs), n = {4} (qubits), L = {2} (layers), N = {1196} (training samples), 2ⁿ = {16} (statevector dimension).\n\n"
        "Space Complexity: O(2ⁿ + N×n + n×L)\n"
        f"= O({16} + {1196}×{4} + {4}×{2}) = O({16 + 1196*4 + 4*2})\n\n"
        "On real quantum hardware, the parameter-shift rule requires 2 circuit evaluations per parameter per gradient step, "
        "giving O(E × P × N) circuit executions where P = 24 parameters. This project uses backpropagation on the simulator "
        "for computational efficiency."
    )

    # ===== CONCLUSION =====
    add_heading("Conclusion")
    add_para(
        f"This project implemented a complete hybrid quantum-classical machine learning pipeline for electricity demand forecasting. "
        f"A 4-qubit Variational Quantum Circuit with StronglyEntanglingLayers was trained using PennyLane's default.qubit simulator "
        f"and compared against a classical Linear Regression baseline.\n\n"
        f"The experimental results clearly show that Linear Regression (R² = {r2_lr:.4f}) significantly outperforms the VQC "
        f"(R² = {r2_vqc:.4f}) on this task. This is attributable to the limited expressivity of a 4-qubit circuit (24 parameters "
        f"in a 16-dimensional Hilbert space), the predominantly linear nature of autocorrelated lag features, and the single-qubit "
        f"measurement bottleneck.\n\n"
        "The project demonstrates that quantum machine learning is not a universal improvement over classical methods, particularly "
        "at the current NISQ (Noisy Intermediate-Scale Quantum) era. The value of this work lies in establishing a reproducible, "
        "end-to-end hybrid quantum-classical ML pipeline — including data processing, quantum circuit design, training, evaluation, "
        "and web deployment — that can be scaled up as quantum hardware matures.\n\n"
        "Future work could explore: (1) more qubits and layers as hardware allows, (2) data re-uploading circuits for increased "
        "expressivity, (3) quantum kernel methods, and (4) problems with inherently non-classical structure where quantum advantage "
        "is more plausible."
    )

    # ===== REFERENCES =====
    add_heading("References")
    refs = [
        "PennyLane Documentation. Xanadu. https://docs.pennylane.ai",
        "AEP Hourly Energy Consumption Dataset. Kaggle (PJM Interconnection). https://www.kaggle.com/datasets/robikscube/hourly-energy-consumption",
    ]
    for i, ref in enumerate(refs, 1):
        doc.add_paragraph(f"[{i}] {ref}")

    docx_path = os.path.join(BASE_DIR, "Capstone_Report.docx")
    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    print("Generating Capstone documents with REAL metrics from artifacts.joblib...")
    print(f"  VQC:  MAE={mae_vqc:.4f}, RMSE={rmse_vqc:.4f}, R²={r2_vqc:.4f}")
    print(f"  LR:   MAE={mae_lr:.4f}, RMSE={rmse_lr:.4f}, R²={r2_lr:.4f}")
    print()
    generate_pptx()
    generate_docx()
    print("\nAll documents generated successfully!")
